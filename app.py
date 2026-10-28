import streamlit as st
import fitz  # PyMuPDF
from PIL import Image
import io
import os
import re
import time
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from google import genai
from google.genai import types
from pydantic import BaseModel, Field

st.set_page_config(page_title="古文書自動校正", page_icon="📝", layout="centered")

# --- セキュリティ：パスワード認証 ---
correct_password = st.secrets.get("APP_PASSWORD", "")

if correct_password:
    if "authenticated" not in st.session_state:
        st.session_state.authenticated = False

    if not st.session_state.authenticated:
        st.title("🔒 パスワード認証")
        pwd_input = st.text_input("アクセス用パスワードを入力してください", type="password")
        if st.button("ログイン", type="primary"):
            if pwd_input == correct_password:
                st.session_state.authenticated = True
                st.rerun()
            else:
                st.error("パスワードが正しくありません。")
        st.stop()

# --- バックアップファイル定義 ---
BACKUP_EXCEL_PATH = "backup_last_result.xlsx"
BACKUP_WORD_PATH = "backup_last_result.docx"

# --- 1. AI出力データ構造の定義 ---
class CorrectionItem(BaseModel):
    page_section: str = Field(description="該当箇所（例: P1・上段など）")
    line_info: str = Field(description="原本の行位置（例: 3行目）")
    error_type: str = Field(description="エラー区分（誤字、脱字、衍字、文字化け、旧字新字等）")
    status: str = Field(description="『確定（修正）』または『要確認（赤字）』")
    original_text: str = Field(description="元のテキスト")
    corrected_text: str = Field(description="修正後のテキスト")
    reason: str = Field(description="修正理由")

class TextSegment(BaseModel):
    text: str = Field(description="テキスト断片")
    mode: str = Field(description="『normal』(標準)、『fixed』(確定修正)、『uncertain』(要確認・赤字)")

class LineData(BaseModel):
    segments: list[TextSegment] = Field(description="原本PDFの1行に含まれるテキスト断片リスト")

class PageData(BaseModel):
    page_number: int = Field(description="原本PDFのページ番号")
    lines: list[LineData] = Field(description="原本PDFの行ごとのデータ。原本画像で改行されている位置と全く同じ単位で1行ずつ格納する")

class ProofreadResult(BaseModel):
    corrections: list[CorrectionItem]
    pages: list[PageData]

# --- 2. テキスト処理関数 ---
def decode_text_file(raw_bytes: bytes) -> str:
    for enc in ["utf-8", "cp932", "shift_jis", "utf-16"]:
        try:
            return raw_bytes.decode(enc)
        except Exception:
            continue
    return raw_bytes.decode("utf-8", errors="replace")

def extract_text_for_batch(full_text: str, b_start: int, b_end: int, total_pdf_pages: int) -> str:
    page_markers = list(re.finditer(r'(?m)^P(\d+)\b', full_text))
    if len(page_markers) >= 3:
        extracted = []
        for i, m in enumerate(page_markers):
            p_num = int(m.group(1))
            if b_start <= p_num <= b_end:
                s_pos = m.start()
                e_pos = page_markers[i+1].start() if i + 1 < len(page_markers) else len(full_text)
                extracted.append(full_text[s_pos:e_pos])
        if extracted:
            return "\n".join(extracted)
    
    blocks = re.split(r'\n\s*\n\s*\n*', full_text.strip())
    if len(blocks) >= total_pdf_pages * 0.5:
        s_idx = max(0, b_start - 1)
        e_idx = min(len(blocks), b_end + 1)
        return "\n\n".join(blocks[s_idx:e_idx])
    
    total_len = len(full_text)
    char_per_page = max(200, total_len / max(1, total_pdf_pages))
    s_char = int(max(0, (b_start - 1.5) * char_per_page))
    e_char = int(min(total_len, (b_end + 0.5) * char_per_page))
    return full_text[s_char:e_char]

# --- 3. Excel・Word生成関数 ---
def create_excel_bytes(corrections: list[CorrectionItem], doc_title: str) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "校正一覧"
    font_family = "Meiryo"
    hf = Font(name=font_family, size=11, bold=True, color="FFFFFF")
    hfill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    tf = Font(name=font_family, size=14, bold=True, color="1F4E78")
    df = Font(name=font_family, size=10)
    bf = Font(name=font_family, size=10, bold=True)
    rf = Font(name=font_family, size=10, bold=True, color="C00000")
    of = Font(name=font_family, size=10, bold=True, color="1E6B34")
    zfill = PatternFill(start_color="F2F6FA", end_color="F2F6FA", fill_type="solid")
    wfill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    afill = PatternFill(start_color="FDE9E9", end_color="FDE9E9", fill_type="solid")
    okfill = PatternFill(start_color="E8F4EA", end_color="E8F4EA", fill_type="solid")
    side = Side(border_style="thin", color="CCCCCC")
    border = Border(left=side, right=side, top=side, bottom=side)

    ws["A1"] = f"{doc_title} 自動校正一覧表"
    ws["A1"].font = tf
    headers = ["No.", "頁・段", "原本行位置", "エラー区分", "ステータス", "修正前", "修正後（正・推定）", "備考・理由"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=c, value=h)
        cell.font, cell.fill, cell.border = hf, hfill, border
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[3].height = 28

    for idx, item in enumerate(corrections, 1):
        r = 3 + idx
        is_alert = "要確認" in item.status
        r_fill = zfill if idx % 2 == 0 else wfill
        vals = [idx, item.page_section, item.line_info, item.error_type, item.status, item.original_text, item.corrected_text, item.reason]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border, cell.fill, cell.font = border, r_fill, df
            if c in [1, 2, 3, 4]: cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c == 5:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.fill, cell.font = (afill, rf) if is_alert else (okfill, of)
            elif c == 7:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                cell.font = rf if is_alert else bf
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[r].height = 36
    for c, w in enumerate([6, 14, 14, 18, 16, 28, 30, 50], 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A4"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()

def create_word_bytes(pages_data: list[PageData], doc_title: str) -> bytes:
    doc = docx.Document()
    style = doc.styles['Normal']
    style.font.name = 'MS Mincho'
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn('w:eastAsia'), 'ＭＳ 明朝')
    
    tp = doc.add_paragraph()
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = tp.add_run(f"{doc_title} 校正済テキスト")
    tr.bold, tr.font.size = True, Pt(15)
    
    lp = doc.add_paragraph()
    lp.paragraph_format.space_after = Pt(6)
    lp.add_run("【凡例】 ").bold = True
    f_run = lp.add_run("紺色太字（下線付）：確定修正")
    f_run.bold, f_run.underline, f_run.font.color.rgb = True, True, RGBColor(0x00, 0x33, 0x66)
    lp.add_run(" ／ ")
    rr = lp.add_run("赤字太字：要確認箇所（虫損・異体字疑義等）")
    rr.bold, rr.font.color.rgb = True, RGBColor(0xC0, 0x00, 0x00)
    
    doc.add_paragraph("━" * 45)

    for page in pages_data:
        hp = doc.add_paragraph()
        hp.paragraph_format.space_before = Pt(14)
        hp.paragraph_format.space_after = Pt(4)
        hr = hp.add_run(f"【 第 {page.page_number} 頁 】")
        hr.bold = True
        hr.font.size = Pt(12)
        hr.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)
        
        for line in page.lines:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.25
            for seg in line.segments:
                r = p.add_run(seg.text)
                r.font.name = 'MS Mincho'
                r._element.rPr.rFonts.set(qn('w:eastAsia'), 'ＭＳ 明朝')
                r.font.size = Pt(11)
                if seg.mode == 'fixed':
                    r.bold, r.underline, r.font.color.rgb = True, True, RGBColor(0x00, 0x33, 0x66)
                elif seg.mode == 'uncertain':
                    r.bold, r.font.color.rgb = True, RGBColor(0xC0, 0x00, 0x00)

        sep_p = doc.add_paragraph()
        sep_p.paragraph_format.space_before = Pt(6)
        sep_p.paragraph_format.space_after = Pt(10)
        sep_r = sep_p.add_run("―" * 40)
        sep_r.font.color.rgb = RGBColor(0xBD, 0xBD, 0xBD)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()

# --- 4. Streamlit UI (本体) ---
st.title("📝 PDF原本×テキスト 自動校正")
st.caption("原本改行完全一致・旧字新字変換・自動救済バックアップ対応。")

# --- セッション切れ救済エリア ---
if os.path.exists(BACKUP_EXCEL_PATH) and os.path.exists(BACKUP_WORD_PATH):
    with st.expander("📦 【直前の校正バックアップ】前回のデータが残っています", expanded=True):
        st.info("通信切断やタイムアウト後でも、直前に完了した成果物をここから再ダウンロードできます。")
        c_b1, c_b2 = st.columns(2)
        with c_b1:
            with open(BACKUP_EXCEL_PATH, "rb") as f:
                st.download_button("📊 直前の校正Excelを復元", f.read(), "直前バックアップ_校正一覧.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", use_container_width=True)
        with c_b2:
            with open(BACKUP_WORD_PATH, "rb") as f:
                st.download_button("📄 直前の校正Wordを復元", f.read(), "直前バックアップ_校正済.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", use_container_width=True)

default_key = st.secrets.get("GEMINI_API_KEY", "") if hasattr(st, "secrets") else ""
api_key = st.sidebar.text_input("Gemini APIキー", value=default_key, type="password")
model_choice = st.sidebar.radio(
    "AIモデル（料金と精度のバランス）", 
    ["1頁あたり約1円 (Flash)", "1頁あたり約4円・高精度 (Pro)"],
    index=0
)
model_name = "gemini-3.8-flash" if "Flash" in model_choice else "gemini-3.8-pro"
convert_kanji = st.sidebar.checkbox("旧漢字を新漢字に変換する", value=True)
batch_size = st.sidebar.slider("1回あたりの処理ページ数", 1, 5, 3)

if "all_corrections" not in st.session_state: st.session_state.all_corrections = []
if "all_pages" not in st.session_state: st.session_state.all_pages = []
if "last_done_page" not in st.session_state: st.session_state.last_done_page = 0

pdf_file = st.file_uploader("① 原本PDF", type=["pdf"])
txt_file = st.file_uploader("② テキストデータ(.txt)", type=["txt"])

if pdf_file and txt_file:
    pdf_bytes = pdf_file.read()
    doc_pdf = fitz.open(stream=pdf_bytes, filetype="pdf")
    total_pages = len(doc_pdf)
    raw_text = decode_text_file(txt_file.read())
    base_name = txt_file.name.rsplit(".", 1)[0]
    st.info(f"PDF読込完了：全 {total_pages} ページ")

    c1, c2 = st.columns(2)
    with c1: start_p = st.number_input("開始ページ", 1, total_pages, 1)
    with c2: end_p = st.number_input("終了ページ", start_p, total_pages, min(start_p + 4, total_pages))

    cb1, cb2 = st.columns([3, 1])
    with cb1: run_btn = st.button("🚀 自動校正スタート", type="primary", use_container_width=True)
    with cb2: 
        if st.button("🗑️ リセット", use_container_width=True):
            st.session_state.all_corrections = []
            st.session_state.all_pages = []
            st.session_state.last_done_page = 0
            st.rerun()

    if run_btn:
        if not api_key:
            st.error("サイドバーからGemini APIキーを入力してください。")
        else:
            client = genai.Client(api_key=api_key)
            st.session_state.all_corrections = []
            st.session_state.all_pages = []
            tot_target = end_p - start_p + 1
            p_bar = st.progress(0, "処理開始...")
            s_box = st.empty()

            kanji_instruction = """
4. **旧字体の新字体化（旧漢字→新漢字）**:
   - 原本画像やテキストにある旧漢字（例: 佛→仏、經→経、體→体、爲→為、應→応、觀→観、釋→釈、晝→昼 等）は、すべて現代の新字体（常用漢字・標準字体）へ変換してWordのテキストに出力してください。
   - 一般的な旧字新字変換はWord本文内で自然に適用し、Excel表にはOCR誤字・脱字や文脈上の疑義箇所を中心に記録してください。
""" if convert_kanji else """
4. **字体の維持**:
   - 原本の字体を尊重し、旧字体を新字体へ意図的に統一・変換することは行わないでください。
"""

            for b_start in range(start_p, end_p + 1, batch_size):
                b_end = min(b_start + batch_size - 1, end_p)
                s_box.info(f"🔄 {b_start}〜{b_end}ページ目を照合中... (対象: {start_p}〜{end_p})")
                imgs = [Image.open(io.BytesIO(doc_pdf[p].get_pixmap(dpi=220).tobytes("png"))) for p in range(b_start-1, b_end)]
                chunk = extract_text_for_batch(raw_text, b_start, b_end, total_pages)
                
                prompt = f"""
あなたは古典籍・仏教文献の専門校正AIです。原本PDF画像({b_start}〜{b_end}頁)と以下のテキストデータを照合・校正してください。

【重要指示】
1. **原本PDFの改行位置の完全一致**:
   - 原本画像で改行されている位置をそのまま維持し、原本の1行を `LineData` の1要素として出力してください。勝手に改行を繋げないでください。
2. **ページ区切り**:
   - 各ページごとに `page_number` を付与し、そのページの行を順番通り漏れなく出力してください。
3. **校正とマーキング**:
   - 明らかな誤字脱字・衍字・文字化け(?など)の修正箇所は statusを『確定（修正）』, modeを『fixed』。
   - 虫損・かすれ・割注脱落など要確認箇所は statusを『要確認（赤字）』, modeを『uncertain』とし、Excel用の理由を記載。
   - 修正のない正常なテキストは modeを『normal』。
{kanji_instruction}

【対象テキスト】\n{chunk}
"""
                for att in range(3):
                    try:
                        res = client.models.generate_content(
                            model=model_name, contents=[*imgs, prompt],
                            config=types.GenerateContentConfig(response_mime_type="application/json", response_schema=ProofreadResult, temperature=0.1)
                        )
                        obj = ProofreadResult.model_validate_json(res.text)
                        st.session_state.all_corrections.extend(obj.corrections)
                        st.session_state.all_pages.extend(obj.pages)
                        st.session_state.last_done_page = b_end
                        break
                    except Exception as e:
                        if att == 2: st.warning(f"⚠️ {b_start}〜{b_end}頁でエラー: {e}")
                        else: time.sleep(5)
                
                # バッチ完了ごとにサーバー内部へ自動バックアップ保存
                if st.session_state.all_corrections or st.session_state.all_pages:
                    tmp_ex = create_excel_bytes(st.session_state.all_corrections, base_name)
                    tmp_wd = create_word_bytes(st.session_state.all_pages, base_name)
                    with open(BACKUP_EXCEL_PATH, "wb") as f: f.write(tmp_ex)
                    with open(BACKUP_WORD_PATH, "wb") as f: f.write(tmp_wd)

                done = b_end - start_p + 1
                p_bar.progress(min(1.0, done / tot_target), f"進行状況: {done} / {tot_target} ページ完了")
            s_box.success(f"✅ {start_p}〜{st.session_state.last_done_page} ページの校正が完了！（検出: {len(st.session_state.all_corrections)} 件）")

    if st.session_state.all_corrections or st.session_state.all_pages:
        st.divider()
        st.subheader(f"📥 ダウンロード（〜P{st.session_state.last_done_page}反映済）")
        ex_data = create_excel_bytes(st.session_state.all_corrections, base_name)
        wd_data = create_word_bytes(st.session_state.all_pages, base_name)
        st.download_button("📊 校正一覧Excelをダウンロード", ex_data, f"{base_name}_校正一覧_P{start_p}-P{st.session_state.last_done_page}.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", use_container_width=True)
        st.download_button("📄 校正済Wordをダウンロード", wd_data, f"{base_name}_校正済_P{start_p}-P{st.session_state.last_done_page}.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", use_container_width=True)
